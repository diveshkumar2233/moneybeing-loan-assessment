import json
from io import BytesIO
from pathlib import Path

import pytest
from app.core.security import create_access_token
from app.models import Lead, Rule, User
from app.schemas.lead import LeadCreate
from app.schemas.rule import RuleWrite
from app.services import credit_score
from app.services.bre_engine import evaluate_lead
from openpyxl import load_workbook
from sqlalchemy import func, select


def seed(db):
    for data in json.loads(
        (
            Path(__file__).resolve().parents[2] / "database/default_rules.json"
        ).read_text()
    ):
        db.add(Rule(**RuleWrite.model_validate(data).model_dump()))
    db.commit()


def test_dynamic_rules_percentage_boundaries(db, payload):
    seed(db)
    application = LeadCreate(**payload)
    assert evaluate_lead(db, application, 742)[0] == "Eligible"
    rule = db.scalar(select(Rule).where(Rule.field == "monthly_income"))
    rule.value = 80000
    db.commit()
    assert (
        "Monthly Income below eligibility criteria"
        in evaluate_lead(db, application, 742)[1]
    )
    rule.is_active = False
    db.commit()
    assert evaluate_lead(db, application, 742)[0] == "Eligible"
    payload["loan_amount"] = "4800000"
    assert evaluate_lead(db, LeadCreate(**payload), 700)[0] == "Eligible"
    payload["loan_amount"] = "4800000.01"
    assert (
        "Loan Amount exceeds eligible limit"
        in evaluate_lead(db, LeadCreate(**payload), 700)[1]
    )


def test_create_duplicate_auth_list_export(client, db, payload, admin, monkeypatch):
    seed(db)
    monkeypatch.setattr(credit_score, "fetch_credit_score", lambda *args: 742)
    response = client.post("/api/leads", json=payload)
    assert response.status_code == 201
    result = response.json()
    assert result == {
        "status": "success",
        "lead_id": result["lead_id"],
        "credit_score": 742,
        "bre_status": "Eligible",
    }
    assert client.post("/api/leads", json=payload).status_code == 409
    assert (
        client.post("/api/leads", json=payload).json()["detail"]
        == "Lead already exists"
    )
    assert db.scalar(select(func.count()).select_from(Lead)) == 1
    assert client.get("/api/leads").status_code == 401
    assert client.get("/api/leads/1").status_code == 401
    response = client.get(
        "/api/leads?search=Asha&loan_type=Home%20Loan&bre_status=Eligible&page_size=1",
        headers=admin,
    )
    assert response.json()["total"] == 1
    assert len(response.json()["items"]) == 1
    assert client.get("/api/leads?search=Nobody", headers=admin).json()["total"] == 0
    summary = client.get("/api/dashboard/summary", headers=admin).json()
    assert summary == {
        "total_leads": 1,
        "eligible_leads": 1,
        "rejected_leads": 0,
        "average_credit_score": 742.0,
    }
    excel = client.get("/api/leads/export", headers=admin)
    workbook = load_workbook(BytesIO(excel.content))
    assert workbook.active.cell(2, 2).value == "Asha Sharma"


def test_credit_failure_is_saved_and_explained(client, db, payload, monkeypatch):
    seed(db)

    def fail(*args):
        raise TimeoutError("provider unavailable")

    monkeypatch.setattr(credit_score, "fetch_credit_score", fail)
    response = client.post("/api/leads", json=payload)
    assert response.status_code == 201
    assert response.json()["credit_score"] is None
    assert response.json()["bre_status"] == "Not Eligible"
    assert "manual review" in response.headers["X-Credit-Score-Error"]
    assert "manual review" in db.scalar(select(Lead)).credit_score_error


@pytest.mark.parametrize(
    "field,value",
    [
        ("mobile", "123"),
        ("pincode", "000000"),
        ("email", "invalid"),
        ("date_of_birth", "2999-01-01"),
        ("date_of_birth", "1800-01-01"),
        ("consent", False),
        ("consent", "true"),
        ("monthly_income", 0),
        ("loan_amount", -1),
        ("property_value", "nan"),
        ("full_name", " "),
        ("loan_type", "Personal Loan"),
    ],
)
def test_validation(client, payload, field, value):
    payload[field] = value
    assert client.post("/api/leads", json=payload).status_code == 422


def test_rule_crud_and_login(client, db, admin):
    login = client.post(
        "/api/auth/login",
        data={"username": "admin@example.com", "password": "TestingPassword123"},
    )
    assert login.status_code == 200
    assert (
        client.post(
            "/api/auth/login",
            data={"username": "admin@example.com", "password": "wrong"},
        ).status_code
        == 401
    )
    data = {
        "rule_name": "Income check",
        "field": "monthly_income",
        "operator": ">=",
        "value": 50000,
        "rejection_message": "Income too low",
        "is_active": True,
    }
    assert client.post("/api/rules", json=data).status_code == 401
    response = client.post("/api/rules", json=data, headers=admin)
    assert response.status_code == 201
    rule_id = response.json()["id"]
    assert client.get(f"/api/rules/{rule_id}", headers=admin).status_code == 200
    data["value"] = 60000
    assert (
        client.put(f"/api/rules/{rule_id}", json=data, headers=admin).json()["value"]
        == "60000.00"
    )
    data["field"] = "__import__"
    assert client.post("/api/rules", json=data, headers=admin).status_code == 422
    assert client.delete(f"/api/rules/{rule_id}", headers=admin).status_code == 204
    assert client.get(f"/api/rules/{rule_id}", headers=admin).status_code == 404
    user = User(email="viewer@example.com", password_hash="unused", role="viewer")
    db.add(user)
    db.commit()
    assert (
        client.get(
            "/api/rules",
            headers={"Authorization": f"Bearer {create_access_token(user.id)}"},
        ).status_code
        == 403
    )


def test_mock_is_repeatable_and_in_range():
    value = credit_score.fetch_credit_score("9876543210", "1995-05-12")
    assert 300 <= value <= 900
    assert value == credit_score.fetch_credit_score("9876543210", "1995-05-12")
